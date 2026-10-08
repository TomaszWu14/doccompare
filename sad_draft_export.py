"""Draft SAD dla agencji celnej: agregacja zatwierdzonych pozycji batcha per
(eksporter × kod CN) — jak pozycje zgłoszenia w WinSAD — i budowa Excela
(arkusze Nagłówek + Pozycje). openpyxl importowany lazy (konwencja repo).
Wzorzec pól: wydruk WinSAD „Podgląd danych zgłoszenia" (spec 2026-09-14)."""

from dataclasses import dataclass, field
from decimal import Decimal

import material_master as mm
from constants import DocKind, InvoiceJobStatus, INVOICE_LIKE_KINDS
from normalizer import normalize_number

SAD_COLUMNS = ["Lp", "Eksporter", "Kod CN", "Kod dod.", "Opis PL", "Ilość SZT",
               "Masa netto", "Kartony", "Wartość [USD]", "Kraj poch.",
               "Nr faktur", "Nr proform", "SENT", "Uwagi"]

_XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _dec(v):
    if v in (None, ""):
        return None
    return normalize_number(str(v))


def _add(acc, d):
    return (acc if acc is not None else Decimal(0)) + d


@dataclass
class _CnGroup:
    """Akumulator jednej pozycji SAD (eksporter × kod CN) w trakcie agregacji."""
    exporter: str
    tariff_cn: str
    customs_code: str = ""
    names: list = field(default_factory=list)
    qty: Decimal | None = None
    qty_note: bool = False
    weight_net: Decimal | None = None
    cartons: Decimal | None = None
    value: Decimal | None = None
    sent: int = 0
    invoices: set = field(default_factory=set)
    proformas: set = field(default_factory=set)
    has_samples: bool = False
    price_fallback: bool = False


def _is_samples_proforma(invoice_number: str, ci_numbers: set) -> bool:
    """Spec 2026-09-14: proforma-próbka = numer kończy się na -S, albo na S gdy
    numer bazowy (bez końcowego S) występuje w batchu jako CI (nie-proforma)."""
    n = (invoice_number or "").strip().upper()
    return n.endswith("-S") or (n.endswith("S") and n[:-1] in ci_numbers)


def aggregate_positions(db, jobs_with_items: list) -> list:
    """jobs_with_items: [{'job': dict, 'items': list[dict]}] (tylko confirmed).
    Wiersz per (supplier_code, tariff_cn); pozycje bez CN → osobny wiersz z uwagą."""
    ci_numbers = {(jw["job"].get("invoice_number") or "").strip().upper()
                  for jw in jobs_with_items
                  if (jw["job"].get("doc_kind") or "") != DocKind.PROFORMA}
    groups = {}
    for jw in jobs_with_items:
        job, items = jw["job"], jw["items"]
        is_proforma = (job.get("doc_kind") or "") == DocKind.PROFORMA
        doc_no = (job.get("invoice_number") or "").strip() or (job.get("filename") or "")
        is_samples = is_proforma and _is_samples_proforma(job.get("invoice_number"), ci_numbers)
        for it in items:
            if int(it.get("skipped") or 0):
                continue
            cn = (it.get("tariff_cn") or "").strip()
            key = (job.get("supplier_code") or "", cn)
            g = groups.setdefault(key, _CnGroup(exporter=key[0], tariff_cn=cn))
            qty = _dec(it.get("qty"))
            factor = _dec(it.get("uom_factor"))
            if qty is not None:
                g.qty = _add(g.qty, qty * (factor if factor else Decimal(1)))
                if not factor and (it.get("uom_src") or "").upper() not in ("", "PCS", "SZT"):
                    g.qty_note = True
            for attr, col in (("weight_net", "weight_net"), ("cartons", "cartons"),
                              ("value", "amount")):
                d = _dec(it.get(col))
                if d is not None:
                    setattr(g, attr, _add(getattr(g, attr), d))
            if not (it.get("net_amount") or "").strip() and (it.get("amount") or "").strip():
                g.price_fallback = True   # ekstraktor bez wartości netto wpisał cenę jedn.
            name = (it.get("name_pl") or "").strip()
            if name and name not in g.names:
                g.names.append(name)
            g.sent = max(g.sent, int(it.get("sent") or 0))
            if is_proforma:
                g.proformas.add(doc_no)
                if is_samples:
                    g.has_samples = True
            else:
                g.invoices.add(doc_no)
            if cn and not g.customs_code and it.get("master_ref"):
                mat = mm.get_material(db, it["master_ref"]) or {}
                g.customs_code = mat.get("customs_code") or ""
    rows = []
    for i, key in enumerate(sorted(groups), start=1):
        g = groups[key]
        opis = "; ".join(g.names)
        if g.has_samples:
            opis = (opis + " WRAZ Z PRÓBKAMI").strip()
        uwagi = []
        if not g.tariff_cn:
            uwagi.append("BRAK CN — uzupełnij master data")
        if g.qty_note:
            uwagi.append("ilość zawiera jednostki źródłowe (brak przelicznika)")
        if g.price_fallback:
            uwagi.append("wartość częściowo z ceny jednostkowej (brak wartości netto pozycji)")
        rows.append({
            "lp": i, "exporter": g.exporter, "tariff_cn": g.tariff_cn,
            "customs_code": g.customs_code, "opis_pl": opis,
            "qty_szt": g.qty, "weight_net": g.weight_net,
            "cartons": g.cartons, "value": g.value,
            "country": "CN",   # v1: kraj wysyłki (spec — rewizja przy dostawcy spoza Chin)
            "invoice_numbers": ", ".join(sorted(g.invoices)),
            "proforma_numbers": ", ".join(sorted(g.proformas)),
            "sent": g.sent, "uwagi": "; ".join(uwagi),
        })
    return rows


def collect_batch(db, batch_id: str):
    """(header, positions) batcha. ValueError, gdy nie wszystkie faktury/proformy
    zatwierdzone — niepełny draft = błędny SAD."""
    import invoice_jobs as ij
    docs = [j for j in ij.list_invoice_jobs(db, batch_id)
            if (j.get("doc_kind") or "invoice") in INVOICE_LIKE_KINDS]
    if not docs:
        raise ValueError("Brak faktur w partii")
    confirmed = [j for j in docs if j.get("status") == InvoiceJobStatus.CONFIRMED]
    if len(confirmed) != len(docs):
        raise ValueError(f"Zatwierdzono {len(confirmed)}/{len(docs)} dokumentów — "
                         "dokończ przegląd przed generowaniem draftu SAD")
    jw = [{"job": j, "items": ij.get_items(db, j["id"])} for j in confirmed]
    positions = aggregate_positions(db, jw)

    def _total(field):
        vals = [r[field] for r in positions if r[field] is not None]
        return sum(vals, Decimal(0)) if vals else None

    gross = None
    for e in jw:
        for it in e["items"]:
            if int(it.get("skipped") or 0):
                continue
            d = _dec(it.get("weight_gross"))
            if d is not None:
                gross = _add(gross, d)
    containers = sorted({(j.get("container_no") or "").strip() for j in confirmed} - {""})
    terms = sorted({(j.get("delivery_terms") or "").strip() for j in confirmed} - {""})
    header = {
        "container": ", ".join(containers),
        "container_warning": len(containers) > 1,
        "delivery_terms": ", ".join(terms),
        "country_dispatch": "CN",
        "currency": "USD",
        "total_value": _total("value"),
        "total_net": _total("weight_net"),
        "total_gross": gross,
        "total_cartons": _total("cartons"),
        "documents": [{"number": (j.get("invoice_number") or j.get("filename") or ""),
                       "supplier": j.get("supplier_code") or "",
                       "kind": j.get("doc_kind") or "invoice",
                       "status": j.get("status") or ""} for j in confirmed],
    }
    return header, positions


def _num(v):
    return "" if v is None else float(v)


def build_sad_workbook(header: dict, positions: list):
    import openpyxl
    from openpyxl.styles import Font
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Nagłówek"
    rows = [
        ("Kontener", header.get("container") or ""),
        ("Warunki dostawy", header.get("delivery_terms") or ""),
        ("Kraj wysyłki", header.get("country_dispatch") or ""),
        ("Waluta", header.get("currency") or ""),
        ("Wartość faktur razem", _num(header.get("total_value"))),
        ("Masa netto razem", _num(header.get("total_net"))),
        ("Masa brutto razem", _num(header.get("total_gross"))),
        ("Liczba kartonów razem", _num(header.get("total_cartons"))),
    ]
    if header.get("container_warning"):
        rows.insert(1, ("UWAGA", "Różne numery kontenerów w partii!"))
    for k, v in rows:
        ws.append([k, v])
    ws.append([])
    ws.append(["Dokumenty:"])
    for d in header.get("documents", []):
        ws.append([d.get("kind", ""), d.get("number", ""), d.get("supplier", ""),
                   d.get("status", "")])
    for c in ws["A"]:
        c.font = Font(bold=True)
    wp = wb.create_sheet("Pozycje")
    wp.append(SAD_COLUMNS)
    for c in wp[1]:
        c.font = Font(bold=True)
    for r in positions:
        wp.append([r["lp"], r["exporter"], r["tariff_cn"], r["customs_code"],
                   r["opis_pl"], _num(r["qty_szt"]), _num(r["weight_net"]),
                   _num(r["cartons"]), _num(r["value"]), r["country"],
                   r["invoice_numbers"], r["proforma_numbers"],
                   "TAK" if r["sent"] else "NIE", r["uwagi"]])
    return wb
