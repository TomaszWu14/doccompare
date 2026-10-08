"""
export_engine.py — eksport raportów do PDF i Excel
Funkcje:
  export_pdf(report_dict, comparison_id) -> bytes
  export_excel(report_dict, comparison_id) -> bytes
"""

import html as _html
import io
import json
import logging
import os
import re as _re
import sys
from datetime import datetime

logger = logging.getLogger(__name__)

MAX_ITEMS_PER_PDF_PAGE = 60
MAX_ITEMS_PER_EXCEL = 5000   # bezpieczny limit wierszy (chroni przed ogromnym workbookiem)

RISK_BG_COLORS = {
    'ok':          '#1a7a4a',
    'ostrzezenie': '#b45309',
    'blad':        '#b91c1c',
}


def _has_keyword(text: str, keyword: str) -> bool:
    """Word-boundary keyword match — avoids 'price' matching 'surprised'."""
    return bool(_re.search(r'\b' + _re.escape(keyword) + r'\b', text, _re.IGNORECASE))

PDF_FONT      = 'Helvetica'
PDF_FONT_BOLD = 'Helvetica-Bold'

# ── PDF via ReportLab ─────────────────────────────────────────────────────────
try:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import (SimpleDocTemplate, Table, TableStyle,
                                     Paragraph, Spacer, HRFlowable)
    from reportlab.lib.enums import TA_CENTER
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfbase import pdfmetrics as _pm

    # Fonty z pełnym wsparciem Unicode (polskie znaki ą ę ó ź ż ć ś ń ł Ł)
    # Próbujemy kilka ścieżek — pierwsza istniejąca wygrywa.
    _FONT_CANDIDATES = [
        # DejaVu (Debian/Ubuntu/Raspbian)
        ('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf',
         '/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf'),
        # Liberation (Red Hat / CentOS / Fedora)
        ('/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf',
         '/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf'),
        # FreeSans (freefont)
        ('/usr/share/fonts/truetype/freefont/FreeSans.ttf',
         '/usr/share/fonts/truetype/freefont/FreeSansBold.ttf'),
        # macOS
        ('/System/Library/Fonts/Supplemental/Arial Unicode MS.ttf',
         '/System/Library/Fonts/Supplemental/Arial Bold.ttf'),
        # Windows
        (os.path.join(os.environ.get('WINDIR','C:/Windows'), 'Fonts/arial.ttf'),
         os.path.join(os.environ.get('WINDIR','C:/Windows'), 'Fonts/arialbd.ttf')),
    ]

    def _try_register_fonts():
        # 1) Najlepiej: osobny regular + bold.
        for reg_path, bold_path in _FONT_CANDIDATES:
            if os.path.exists(reg_path) and os.path.exists(bold_path):
                try:
                    _pm.registerFont(TTFont('DocFont',     reg_path))
                    _pm.registerFont(TTFont('DocFont-Bold', bold_path))
                    _pm.registerFontFamily('DocFont',
                                           normal='DocFont',
                                           bold='DocFont-Bold')
                    return True
                except Exception as exc:
                    logger.warning("Font rejestracja nieudana (%s): %s", reg_path, exc)
        # 2) Fallback: jest sam regular (brak wariantu bold) — użyj go również jako
        #    „bold". Lepsze „pogrubienie = zwykły Unicode" niż Helvetica (krzaki ą/ę/ó).
        for reg_path, _bold_path in _FONT_CANDIDATES:
            if os.path.exists(reg_path):
                try:
                    _pm.registerFont(TTFont('DocFont',      reg_path))
                    _pm.registerFont(TTFont('DocFont-Bold', reg_path))
                    _pm.registerFontFamily('DocFont',
                                           normal='DocFont',
                                           bold='DocFont-Bold')
                    return True
                except Exception as exc:
                    logger.warning("Font fallback nieudany (%s): %s", reg_path, exc)
        return False

    if _try_register_fonts():
        PDF_FONT      = 'DocFont'
        PDF_FONT_BOLD = 'DocFont-Bold'
    else:
        logger.warning("Brak fontu Unicode — polskie znaki (ą/ę/ó/ź/ż) mogą renderować "
                       "się jako kwadraty. Zainstaluj: apt install fonts-dejavu-core")

    HAS_REPORTLAB = True
except ImportError:
    HAS_REPORTLAB = False

# ── Excel via openpyxl ────────────────────────────────────────────────────────
try:
    import openpyxl
    from openpyxl.styles import (Font, PatternFill, Alignment, Border, Side)
    from openpyxl.utils import get_column_letter
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

def _safe(v, default=''):
    """Zamień None/nan na pusty string lub default."""
    if v is None: return default
    s = str(v).strip()
    return default if s.lower() in ('none','nan','null','—','') else s


_FORMULA_PREFIXES = ('=', '+', '-', '@', '\t', '\r')

def _xsafe(v, default=''):
    """Jak _safe, ale zabezpiecza przed Excel formula injection (prefix apostrofem)."""
    s = _safe(v, default)
    return ("'" + s) if (s and s[0] in _FORMULA_PREFIXES) else s


def _safe_para(v, default='—'):
    """Like _safe but also escapes HTML entities for use in ReportLab Paragraph."""
    return _html.escape(_safe(v, default))



# ─────────────────────────────────────────────────────────────────────────────
# KOLORY — definiowane tylko gdy reportlab jest dostępny, ponieważ `colors`
# jest importowany warunkowo. W przeciwnym razie moduł nie załaduje się wcale.
# ─────────────────────────────────────────────────────────────────────────────
if HAS_REPORTLAB:
    C_ERR   = colors.HexColor('#f74f4f')
    C_WARN  = colors.HexColor('#f7a84f')
    C_OK    = colors.HexColor('#4fce8e')
    C_FMT   = colors.HexColor('#7b9ef7')
    C_DARK  = colors.HexColor('#1e2130')
    C_MUTED = colors.HexColor('#6b7080')
    C_WHITE = colors.white
    C_LIGHT = colors.HexColor('#f8f9fc')
    C_ACCENT= colors.HexColor('#4f8ef7')


# ─────────────────────────────────────────────────────────────────────────────
# PDF STYLE CACHE
# ─────────────────────────────────────────────────────────────────────────────

_PDF_STYLES = None


def _get_pdf_styles():
    """Return cached getSampleStyleSheet() instance (created once, reused per process)."""
    global _PDF_STYLES
    if _PDF_STYLES is None:
        _PDF_STYLES = getSampleStyleSheet()
    return _PDF_STYLES


# ─────────────────────────────────────────────────────────────────────────────
# PDF EXPORT
# ─────────────────────────────────────────────────────────────────────────────

def export_pdf(report: dict, comparison_id: int = None) -> bytes:
    """Generuje raport PDF. Zwraca bytes."""
    if not HAS_REPORTLAB:
        raise ImportError("reportlab nie jest zainstalowany. pip install reportlab")

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4,
                            leftMargin=1.5*cm, rightMargin=1.5*cm,
                            topMargin=1.5*cm, bottomMargin=1.5*cm)

    styles = _get_pdf_styles()
    story = []

    def S(name, parent='Normal', **kw):
        return ParagraphStyle(name, parent=styles[parent], **kw)

    sSection = S('sec', fontSize=8.5, textColor=C_MUTED, spaceBefore=10, spaceAfter=4,
                 fontName=PDF_FONT_BOLD, leading=12)
    sSmall   = S('sm', fontSize=7.5, textColor=C_MUTED, leading=10, fontName=PDF_FONT)
    sCell    = S('cell', fontSize=8, leading=10, fontName=PDF_FONT)
    sCellB   = S('cellb', fontSize=8, leading=10, fontName=PDF_FONT_BOLD)
    # Nagłówek tabeli: BIAŁY tekst (wiersz nagłówka ma ciemne tło #1e293b).
    # Kolor Paragraph wygrywa z TEXTCOLOR tabeli, więc musi być tu jawnie biały.
    sCellHdr = S('cellhdr', fontSize=8, leading=10, fontName=PDF_FONT_BOLD, textColor=C_WHITE)

    now  = datetime.now().strftime('%d.%m.%Y %H:%M')
    risk = report.get('risk_level', 'ok')
    modules = report.get('modules', {})
    tbl_m   = modules.get('table', {})
    typo_m  = modules.get('typo', {})
    rx_m    = modules.get('regex', {})

    fa  = _safe(report.get('file_a'), '—')
    fb  = _safe(report.get('file_b'), '—')
    ta  = report.get('doc_type_a', '')
    tb  = report.get('doc_type_b', '')
    W   = 17.7*cm   # usable width

    # ── 1. PASEK NAGŁÓWKA ────────────────────────────────────────────────────
    risk_bg    = {'ok': '#1a7a4a', 'ostrzezenie': '#b45309', 'blad': '#b91c1c'}.get(risk, '#374151')
    _risk_labels = {'ok': 'ZGODNE', 'ostrzezenie': 'WYMAGA WERYFIKACJI', 'blad': 'BŁĘDY KRYTYCZNE'}
    risk_label = _risk_labels.get(risk) or _html.escape(str(risk or "").upper())

    hdr_tbl = Table([[
        Paragraph(f'<font size="13"><b>DocCompare</b></font>  '
                  f'<font size="9" color="#a5b4fc">Raport porównania dokumentów</font>',
                  S('ht', fontSize=9, textColor=C_WHITE, fontName=PDF_FONT, leading=16)),
        Paragraph(f'<b>{risk_label}</b>',
                  S('hs', fontSize=9, textColor=C_WHITE, fontName=PDF_FONT_BOLD,
                    alignment=TA_CENTER, leading=12)),
        Paragraph(f'#{comparison_id or "—"}<br/><font size="7" color="#a5b4fc">{now}</font>',
                  S('hd', fontSize=8, textColor=C_WHITE, fontName=PDF_FONT,
                    alignment=TA_CENTER, leading=11)),
    ]], colWidths=[9*cm, 5.5*cm, 3.2*cm])
    hdr_tbl.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), colors.HexColor(risk_bg)),
        ('BACKGROUND', (1,0), (1,0),   colors.HexColor('#1e293b')),
        ('TOPPADDING',    (0,0), (-1,-1), 7),
        ('BOTTOMPADDING', (0,0), (-1,-1), 7),
        ('LEFTPADDING',   (0,0), (-1,-1), 10),
        ('RIGHTPADDING',  (0,0), (-1,-1), 10),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
    ]))
    story.append(hdr_tbl)
    story.append(Spacer(1, 0.25*cm))

    # ── 2. KARTY DOKUMENTÓW (PO | VS | PI) ───────────────────────────────────
    def _doc_card_rows(label, filename, doc_type, meta_dict):
        rows = []
        for k, v in meta_dict.items():
            if v and _safe(v):
                rows.append(f'<font color="#64748b" size="7">{_html.escape(str(k))}</font><br/>'
                            f'<font size="8"><b>{_html.escape(_safe(v, "")[:40])}</b></font>')
        cell_content = (
            f'<font size="7" color="#94a3b8">{_html.escape((doc_type or "AUTO").upper())}</font><br/>'
            f'<font size="9"><b>{_html.escape(filename[:35])}</b></font><br/>'
            + '<br/>'.join(rows[:5])
        )
        return cell_content

    headers_list = tbl_m.get('headers', []) if tbl_m.get('ok') else []
    def _hval(key, side):
        for h in headers_list:
            if h.get('key','').lower() == key.lower():
                return _safe(h.get(f'val_{side}'), '')
        return ''

    meta_a = {
        'Numer': _hval('numer_zamowienia', 'a') or _hval('nr', 'a'),
        'Data':  _hval('data', 'a'),
        'Waluta': _hval('waluta', 'a'),
        'Incoterms': _hval('incoterms', 'a'),
        'Płatność': _hval('platnosc', 'a') or _hval('payment', 'a'),
    }
    meta_b = {
        'Numer': _hval('numer_zamowienia', 'b') or _hval('nr', 'b'),
        'Data':  _hval('data', 'b'),
        'Waluta': _hval('waluta', 'b'),
        'Incoterms': _hval('incoterms', 'b'),
        'Płatność': _hval('platnosc', 'b') or _hval('payment', 'b'),
    }

    card_a = _doc_card_rows(ta or 'DOK A', fa, ta, meta_a)
    card_b = _doc_card_rows(tb or 'DOK B', fb, tb, meta_b)
    vs_para = Paragraph('<b>VS</b>', S('vs', fontSize=14, textColor=C_MUTED,
                                       fontName=PDF_FONT_BOLD, alignment=TA_CENTER))

    cards_tbl = Table([
        [Paragraph(card_a, S('ca', fontSize=8, leading=11, fontName=PDF_FONT)),
         vs_para,
         Paragraph(card_b, S('cb', fontSize=8, leading=11, fontName=PDF_FONT))]
    ], colWidths=[7.8*cm, 1.5*cm, 8.4*cm])
    cards_tbl.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (0,0), colors.HexColor('#eff6ff')),
        ('BACKGROUND', (2,0), (2,0), colors.HexColor('#f0fdf4')),
        ('BOX', (0,0), (0,0), 1, colors.HexColor('#3b82f6')),
        ('BOX', (2,0), (2,0), 1, colors.HexColor('#22c55e')),
        ('TOPPADDING',    (0,0), (-1,-1), 8),
        ('BOTTOMPADDING', (0,0), (-1,-1), 8),
        ('LEFTPADDING',   (0,0), (-1,-1), 8),
        ('RIGHTPADDING',  (0,0), (-1,-1), 8),
        ('VALIGN', (0,0), (-1,-1), 'TOP'),
        ('VALIGN', (1,0), (1,0), 'MIDDLE'),
    ]))
    story.append(cards_tbl)
    story.append(Spacer(1, 0.25*cm))

    # ── 3. KPI TILES ─────────────────────────────────────────────────────────
    items_all  = tbl_m.get('items', []) if tbl_m.get('ok') else []
    n_total    = len(items_all)
    n_ok       = sum(1 for i in items_all if i.get('status') == 'ok')
    n_diff     = sum(1 for i in items_all if i.get('status') in ('roznica','brak_w_a','brak_w_b'))
    total_a    = _safe(tbl_m.get('total_a'), '—')
    total_b    = _safe(tbl_m.get('total_b'), '—')

    try:
        from normalizer import normalize_number as _nn
        _ta = _nn(str(total_a))
        _tb = _nn(str(total_b))
        if _ta is None or _tb is None:
            raise ValueError("unparseable total")
        ta_f, tb_f = float(_ta), float(_tb)
        diff_val = f'{tb_f - ta_f:+.2f}'
    except Exception:
        diff_val = '—'
        ta_f = tb_f = 0.0

    score_pct = round(n_ok / n_total * 100, 1) if n_total else 0.0

    kpi_bgs = ['#f8fafc','#f0fdf4','#fef2f2','#eff6ff','#f0fdf4','#fefce8']
    kpi_borders = ['#cbd5e1','#22c55e','#f87171','#3b82f6','#22c55e','#eab308']

    kpi_row = []
    for idx, (label, value, bg) in enumerate([
        ('POZYCJE', str(n_total), '#f8fafc'),
        ('ZGODNE', str(n_ok), '#f0fdf4'),
        ('ROZBIEŻNOŚCI', str(n_diff), '#fef2f2'),
        ('WARTOŚĆ A', _html.escape(str(total_a)), '#eff6ff'),
        ('WARTOŚĆ B', _html.escape(str(total_b)), '#f0fdf4'),
        ('RÓŻNICA', _html.escape(str(diff_val)), '#fefce8'),
    ]):
        kpi_row.append(Paragraph(
            f'<font size="6.5" color="#64748b">{label}</font><br/><b>{value}</b>',
            S(f'k{idx}', fontSize=10, leading=13, fontName=PDF_FONT, alignment=TA_CENTER)))

    kpi_tbl = Table([kpi_row], colWidths=[W/6]*6)
    kpi_cmds = [
        ('TOPPADDING',    (0,0), (-1,-1), 7),
        ('BOTTOMPADDING', (0,0), (-1,-1), 7),
        ('LEFTPADDING',   (0,0), (-1,-1), 4),
        ('RIGHTPADDING',  (0,0), (-1,-1), 4),
        ('VALIGN',  (0,0), (-1,-1), 'MIDDLE'),
        ('FONTSIZE',(0,0), (-1,-1), 10),
        ('FONTNAME',(0,0), (-1,-1), PDF_FONT_BOLD),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#e2e8f0')),
    ]
    for idx, bg in enumerate(kpi_bgs):
        kpi_cmds.append(('BACKGROUND', (idx,0), (idx,0), colors.HexColor(bg)))
        kpi_cmds.append(('LINEBELOW', (idx,0), (idx,0), 2.5, colors.HexColor(kpi_borders[idx])))
    kpi_tbl.setStyle(TableStyle(kpi_cmds))
    story.append(kpi_tbl)
    story.append(Spacer(1, 0.3*cm))

    # ── 4. KLUCZOWE ROZBIEŻNOŚCI ─────────────────────────────────────────────
    div_items = [i for i in items_all if i.get('status') in ('roznica','brak_w_a','brak_w_b')]
    div_hdrs  = [h for h in headers_list
                 if h.get('status') in ('roznica','brak_w_a','brak_w_b')]
    typo_errs = [f for f in (typo_m.get('findings') or []) if f.get('severity') == 'error'] if typo_m.get('ok') else []

    # classify severity — use word-boundary matching to avoid false positives
    # e.g. 'net' should not match 'internet', 'price' should not match 'surprised'
    critical_divs, major_divs, info_divs = [], [], []
    for it in div_items:
        issues = it.get('issues') or []
        issue_str = ' '.join(str(x) for x in issues)
        if _has_keyword(issue_str, 'brak') or it.get('status') in ('brak_w_a','brak_w_b'):
            critical_divs.append(('item', it))
        elif (_has_keyword(issue_str, 'cena') or _has_keyword(issue_str, 'net')
              or _has_keyword(issue_str, 'wartosc')):
            major_divs.append(('item', it))
        else:
            info_divs.append(('item', it))
    for h in div_hdrs:
        if h.get('status') in ('brak_w_a','brak_w_b'):
            critical_divs.append(('hdr', h))
        else:
            major_divs.append(('hdr', h))
    for t in typo_errs:
        info_divs.append(('typo', t))

    all_divs = ([(d, 'KRYTYCZNE', '#fef2f2', '#dc2626') for d in critical_divs[:3]] +
                [(d, 'WIĘKSZE',   '#fff7ed', '#ea580c') for d in major_divs[:4]] +
                [(d, 'INFORMACYJNE', '#f0f9ff', '#0284c7') for d in info_divs[:3]])

    if all_divs:
        story.append(Paragraph('KLUCZOWE ROZBIEŻNOŚCI', sSection))
        div_rows = []
        row_buf = []
        for (dtype, ddata), sev_label, sev_bg, sev_fg in all_divs:
            if dtype == 'item':
                title = _safe_para(ddata.get('ref'), '—')
                before = _safe_para(ddata.get('net_a'), '—')
                after  = _safe_para(ddata.get('net_b'), '—')
                detail = _html.escape('; '.join(str(x) for x in (ddata.get('issues') or []))[:50])
            elif dtype == 'hdr':
                title = _safe_para(ddata.get('key'), '—')
                before = _safe_para(ddata.get('val_a'), '—')
                after  = _safe_para(ddata.get('val_b'), '—')
                detail = _html.escape((ddata.get('comment') or '')[:50])
            else:
                title = _safe_para(ddata.get('field_name') or ddata.get('category'), '—')
                before = _safe_para(ddata.get('val_a'), '—')
                after  = _safe_para(ddata.get('val_b'), '—')
                detail = _html.escape((ddata.get('description') or '')[:50])

            cell_para = Paragraph(
                f'<font size="6.5" color="{sev_fg}"><b>{sev_label}</b></font>  '
                f'<font size="8"><b>{title}</b></font><br/>'
                f'<font size="7" color="#64748b">przed: </font>'
                f'<font size="7.5"><strike>{before}</strike></font>'
                f'  <font size="7" color="#64748b">po: </font>'
                f'<font size="8"><b>{after}</b></font>'
                + (f'<br/><font size="7" color="#94a3b8">{detail}</font>' if detail else ''),
                S(f'dv_{title}', fontSize=8, leading=11, fontName=PDF_FONT))
            card = Table([[cell_para]], colWidths=[5.4*cm])
            card.setStyle(TableStyle([
                ('BACKGROUND',    (0,0), (-1,-1), colors.HexColor(sev_bg)),
                ('LEFTPADDING',   (0,0), (-1,-1), 7),
                ('RIGHTPADDING',  (0,0), (-1,-1), 7),
                ('TOPPADDING',    (0,0), (-1,-1), 6),
                ('BOTTOMPADDING', (0,0), (-1,-1), 6),
                ('LINEABOVE', (0,0), (-1,0), 2.5, colors.HexColor(sev_fg)),
                ('BOX', (0,0), (-1,-1), 0.4, colors.HexColor('#e2e8f0')),
            ]))
            row_buf.append(card)
            if len(row_buf) == 3:
                div_rows.append(row_buf)
                row_buf = []
        if row_buf:
            while len(row_buf) < 3:
                row_buf.append(Paragraph('', sSmall))
            div_rows.append(row_buf)

        div_grid = Table(div_rows, colWidths=[5.7*cm, 5.7*cm, 5.7*cm],
                         hAlign='LEFT')
        div_grid.setStyle(TableStyle([
            ('LEFTPADDING',   (0,0), (-1,-1), 3),
            ('RIGHTPADDING',  (0,0), (-1,-1), 3),
            ('TOPPADDING',    (0,0), (-1,-1), 3),
            ('BOTTOMPADDING', (0,0), (-1,-1), 3),
            ('VALIGN', (0,0), (-1,-1), 'TOP'),
        ]))
        story.append(div_grid)
        story.append(Spacer(1, 0.3*cm))

    # ── 5. TABELA POZYCJI TOWAROWYCH ─────────────────────────────────────────
    if items_all:
        story.append(Paragraph(
            f'POZYCJE TOWAROWE — {n_total} pozycji  |  '
            f'zgodnych: {n_ok}  |  rozbieżności: {n_diff}',
            sSection))

        rows = [[
            Paragraph('ST', sCellHdr),
            Paragraph('REF', sCellHdr),
            Paragraph('OPIS', sCellHdr),
            Paragraph('Ilość A', sCellHdr),
            Paragraph('Ilość B', sCellHdr),
            Paragraph('NET A', sCellHdr),
            Paragraph('NET B', sCellHdr),
            Paragraph('Δ WARTOŚĆ', sCellHdr),
        ]]
        status_map = {
            'ok': ('✓', C_OK),
            'roznica': ('!', C_ERR),
            'brak_w_a': ('!', C_ERR),
            'brak_w_b': ('!', C_ERR),
            'format': ('~', C_FMT),
            'watpliwe': ('?', C_WARN),
        }
        row_cmds = []
        for ri, it in enumerate(items_all[:60], 1):
            st   = it.get('status', '')
            icon, _ = status_map.get(st, ('·', C_MUTED))
            ref  = _html.escape(_safe(it.get('ref'), '—')[:18])
            desc = _html.escape(_safe(it.get('desc_b') or it.get('desc_a'), '')[:30])
            qa   = _safe_para(it.get('qty_a'), '—')
            qb   = _safe_para(it.get('qty_b'), '—')
            _na_raw = _safe(it.get('net_a'), '—')
            _nb_raw = _safe(it.get('net_b'), '—')
            na   = _html.escape(_na_raw)
            nb   = _html.escape(_nb_raw)
            try:
                from normalizer import normalize_number as _nn_dv
                _na_n = _nn_dv(str(_na_raw))
                _nb_n = _nn_dv(str(_nb_raw))
                if _na_n is None or _nb_n is None:
                    raise ValueError("unparseable net")
                dv = f'{float(_nb_n) - float(_na_n):+.2f}'
            except Exception:
                dv = '—'
            rows.append([
                Paragraph(icon, S(f'si{ri}', fontSize=8, fontName=PDF_FONT_BOLD,
                                  textColor=status_map.get(st,('·',C_MUTED))[1],
                                  alignment=TA_CENTER)),
                Paragraph(ref,  sCell),
                Paragraph(desc, sCell),
                Paragraph(qa,   sCell),
                Paragraph(qb,   sCell),
                Paragraph(na,   sCell),
                Paragraph(nb,   sCell),
                Paragraph(dv,   S(f'dv{ri}', fontSize=8, fontName=PDF_FONT_BOLD,
                                  textColor=(C_ERR if dv.startswith('+') and dv != '+0.00'
                                             else C_OK if dv.startswith('-') else C_MUTED),
                                  alignment=TA_CENTER)),
            ])
            if st in ('roznica','brak_w_a','brak_w_b'):
                row_cmds.append(('BACKGROUND', (0,ri), (-1,ri), colors.HexColor('#fff5f5')))
            elif st == 'format':
                row_cmds.append(('BACKGROUND', (0,ri), (-1,ri), colors.HexColor('#eff6ff')))

        # RAZEM row
        rows.append([
            Paragraph('', sCell),
            Paragraph('RAZEM', S('raz', fontSize=8, fontName=PDF_FONT_BOLD)),
            Paragraph('', sCell),
            Paragraph('', sCell), Paragraph('', sCell),
            Paragraph(_html.escape(str(total_a)), S('ta', fontSize=8, fontName=PDF_FONT_BOLD)),
            Paragraph(_html.escape(str(total_b)), S('tb', fontSize=8, fontName=PDF_FONT_BOLD)),
            Paragraph(_html.escape(diff_val), S('tdiff', fontSize=8, fontName=PDF_FONT_BOLD,
                                  textColor=C_ERR if (diff_val.startswith('+') and diff_val != '+0.00') else C_OK)),
        ])
        last = len(rows) - 1
        row_cmds.append(('BACKGROUND', (0, last), (-1, last), colors.HexColor('#f1f5f9')))
        row_cmds.append(('LINEABOVE',  (0, last), (-1, last), 1, C_MUTED))

        items_t = Table(rows, colWidths=[0.9*cm, 2.8*cm, 4.2*cm, 1.6*cm,
                                         1.6*cm, 2.2*cm, 2.2*cm, 2.2*cm])
        base_cmds = [
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#1e293b')),
            ('TEXTCOLOR',  (0,0), (-1,0), C_WHITE),
            ('FONTNAME',   (0,0), (-1,0), PDF_FONT_BOLD),
            ('FONTSIZE',   (0,0), (-1,-1), 8),
            ('ROWBACKGROUNDS', (0,1), (-1, last-1), [C_WHITE, C_LIGHT]),
            ('GRID',   (0,0), (-1,-1), 0.3, colors.HexColor('#e2e8f0')),
            ('TOPPADDING',    (0,0), (-1,-1), 3),
            ('BOTTOMPADDING', (0,0), (-1,-1), 3),
            ('LEFTPADDING',   (0,0), (-1,-1), 4),
            ('VALIGN', (0,0), (-1,-1), 'TOP'),
        ]
        items_t.setStyle(TableStyle(base_cmds + row_cmds))
        story.append(items_t)
        if len(items_all) > 60:
            story.append(Paragraph(f'… pokazano 60 z {len(items_all)} pozycji. Pełna lista w Excel.',
                                   sSmall))
        story.append(Spacer(1, 0.3*cm))

    # ── 6. TABELA NAGŁÓWKÓW ───────────────────────────────────────────────────
    if headers_list:
        story.append(Paragraph('PORÓWNANIE PÓŁ NAGŁÓWKOWYCH', sSection))
        STATUS_LABELS = {
            'ok':        ('ZGODNE',   '#dcfce7', '#16a34a'),
            'format':    ('SPRAWDŹ',  '#fef9c3', '#ca8a04'),
            'watpliwe':  ('SPRAWDŹ',  '#fef9c3', '#ca8a04'),
            'roznica':   ('NIEZGODNE','#fee2e2', '#dc2626'),
            'brak_w_a':  ('TYLKO B',  '#fce7f3', '#be185d'),
            'brak_w_b':  ('TYLKO A',  '#fce7f3', '#be185d'),
        }
        hdr_rows = [[
            Paragraph('POLE', sCellHdr),
            Paragraph('ZAMÓWIENIE (A)', sCellHdr),
            Paragraph('PROFORMA (B)', sCellHdr),
            Paragraph('STATUS', sCellHdr),
        ]]
        hdr_cmds = []
        for ri, h in enumerate(headers_list[:25], 1):
            st  = h.get('status', 'ok')
            lbl, bg, fg = STATUS_LABELS.get(st, ('—', '#f8fafc', '#64748b'))
            badge = Table([[Paragraph(f'<b>{lbl}</b>',
                                      S(f'b{ri}', fontSize=7, fontName=PDF_FONT_BOLD,
                                        textColor=colors.HexColor(fg),
                                        alignment=TA_CENTER))]],
                          colWidths=[2.5*cm])
            badge.setStyle(TableStyle([
                ('BACKGROUND', (0,0), (-1,-1), colors.HexColor(bg)),
                ('BOX', (0,0), (-1,-1), 0.5, colors.HexColor(fg)),
                ('TOPPADDING',    (0,0), (-1,-1), 2),
                ('BOTTOMPADDING', (0,0), (-1,-1), 2),
                ('LEFTPADDING',   (0,0), (-1,-1), 4),
                ('RIGHTPADDING',  (0,0), (-1,-1), 4),
            ]))
            hdr_rows.append([
                Paragraph(_html.escape(_safe(h.get('key'), '—')[:30]), sCell),
                Paragraph(_html.escape(_safe(h.get('val_a'), '—')[:40]), sCell),
                Paragraph(_html.escape(_safe(h.get('val_b'), '—')[:40]), sCell),
                badge,
            ])
            if st in ('roznica','brak_w_a','brak_w_b'):
                hdr_cmds.append(('BACKGROUND', (0,ri), (2,ri), colors.HexColor('#fff5f5')))

        hdr_t = Table(hdr_rows, colWidths=[3.5*cm, 5.7*cm, 5.7*cm, 2.8*cm])
        hdr_t.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#1e293b')),
            ('TEXTCOLOR',  (0,0), (-1,0), C_WHITE),
            ('FONTNAME',   (0,0), (-1,0), PDF_FONT_BOLD),
            ('FONTSIZE',   (0,0), (-1,-1), 8),
            ('ROWBACKGROUNDS', (0,1), (-1,-1), [C_WHITE, C_LIGHT]),
            ('GRID',   (0,0), (-1,-1), 0.3, colors.HexColor('#e2e8f0')),
            ('TOPPADDING',    (0,0), (-1,-1), 3),
            ('BOTTOMPADDING', (0,0), (-1,-1), 3),
            ('LEFTPADDING',   (0,0), (-1,-1), 5),
            ('VALIGN', (0,0), (-1,-1), 'TOP'),
        ] + hdr_cmds))
        story.append(hdr_t)
        story.append(Spacer(1, 0.3*cm))

    # ── 7. ROZLICZENIE FINANSOWE ──────────────────────────────────────────────
    if total_a != '—' or total_b != '—':
        story.append(Paragraph('ROZLICZENIE FINANSOWE', sSection))
        try:
            from normalizer import normalize_number as _nn2
            _ta2 = _nn2(str(total_a)); _tb2 = _nn2(str(total_b))
            if _ta2 is None or _tb2 is None:
                raise ValueError("unparseable total")
            ta_f, tb_f = float(_ta2), float(_tb2)
            max_v = max(abs(ta_f), abs(tb_f)) or 1.0
            bar_a = int(abs(ta_f) / max_v * 120)
            bar_b = int(abs(tb_f) / max_v * 120)
            diff_abs = abs(tb_f - ta_f)
            diff_pct = diff_abs / ta_f * 100 if ta_f else 0

            fin_data = [
                [Paragraph('Wartość A (Zamówienie)', sCell),
                 Paragraph('▓' * (bar_a // 10) + ' ' + _html.escape(str(total_a)),
                           S('ba', fontSize=8, textColor=colors.HexColor('#3b82f6'),
                             fontName=PDF_FONT_BOLD))],
                [Paragraph('Wartość B (Proforma)', sCell),
                 Paragraph('▓' * (bar_b // 10) + ' ' + _html.escape(str(total_b)),
                           S('bb', fontSize=8, textColor=colors.HexColor('#22c55e'),
                             fontName=PDF_FONT_BOLD))],
                [Paragraph('Różnica', S('rd', fontSize=8, fontName=PDF_FONT_BOLD,
                                        textColor=C_ERR if diff_abs > 0 else C_OK)),
                 Paragraph(_html.escape(diff_val) + f'  ({diff_pct:.1f}%)',
                           S('dif', fontSize=8, fontName=PDF_FONT_BOLD,
                             textColor=C_ERR if diff_abs > 0 else C_OK))],
            ]
            fin_t = Table(fin_data, colWidths=[4.5*cm, 13.2*cm])
            fin_t.setStyle(TableStyle([
                ('FONTSIZE',  (0,0), (-1,-1), 8),
                ('ROWBACKGROUNDS', (0,0), (-1,-1), [C_WHITE, C_LIGHT, colors.HexColor('#fef9c3')]),
                ('GRID', (0,0), (-1,-1), 0.3, colors.HexColor('#e2e8f0')),
                ('TOPPADDING',    (0,0), (-1,-1), 5),
                ('BOTTOMPADDING', (0,0), (-1,-1), 5),
                ('LEFTPADDING',   (0,0), (-1,-1), 6),
                ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
            ]))
            story.append(fin_t)
        except Exception:
            pass
        story.append(Spacer(1, 0.2*cm))

    # ── STOPKA ────────────────────────────────────────────────────────────────
    story.append(HRFlowable(width='100%', thickness=0.5, color=C_MUTED, spaceBefore=8))
    story.append(Paragraph(
        f'Wygenerowano: {now}  |  DocCompare v6  |  ACME sp. z o.o.',
        S('foot', fontSize=7, textColor=C_MUTED, alignment=TA_CENTER)))

    doc.build(story)
    return buf.getvalue()


def _collect_critical(report):
    crits = []
    m = report.get('modules', {})
    if m.get('table', {}).get('ok'):
        for it in m['table'].get('items', []):
            if it.get('status') == 'roznica':
                crits.append(f"[Tabela] {it.get('ref','')}: {'; '.join(str(x) for x in (it.get('issues') or []))}")
        for h in m['table'].get('headers', []):
            if h.get('status') in ('roznica', 'brak_w_a', 'brak_w_b'):
                crits.append(f"[Tabela] '{h.get('key','')}':"
                             f" {h.get('val_a','—')} → {h.get('val_b','—')}")
    if m.get('typo', {}).get('ok'):
        for f in m['typo'].get('findings', []):
            if f.get('severity') == 'error':
                crits.append(f"[I↔1] {f.get('description','')}")
    if m.get('regex', {}).get('ok'):
        for f in m['regex'].get('fields', []):
            if f.get('status') in ('roznica', 'brak_w_a', 'brak_w_b'):
                crits.append(f"[Regex] '{f.get('field','')}': "
                             f"{f.get('doc_a','—')} vs {f.get('doc_b','—')}")
    return crits


def _table_style_cmds(nrows):
    return [
        ('FONTNAME',  (0,0), (-1,0),  PDF_FONT_BOLD),
        ('FONTSIZE',  (0,0), (-1,-1), 8),
        ('BACKGROUND',(0,0), (-1,0),  colors.HexColor('#2a3050')),
        ('TEXTCOLOR', (0,0), (-1,0),  C_WHITE),
        ('ROWBACKGROUNDS', (0,1), (-1,-1), [C_WHITE, C_LIGHT]),
        ('GRID',      (0,0), (-1,-1), 0.3, colors.HexColor('#d0d4e0')),
        ('TOPPADDING', (0,0), (-1,-1), 3),
        ('BOTTOMPADDING', (0,0), (-1,-1), 3),
        ('LEFTPADDING', (0,0), (-1,-1), 5),
        ('VALIGN', (0,0), (-1,-1), 'TOP'),
    ]


def _table_style(nrows):
    return TableStyle(_table_style_cmds(nrows))


def _items_table_style(rows):
    base = _table_style_cmds(len(rows))
    extra = []
    status_col = 5
    color_map = {
        'roznica': C_ERR, 'brak_w_a': C_ERR, 'brak_w_b': C_ERR,
        'format': C_FMT, 'watpliwe': C_WARN, 'ok': C_OK,
    }
    for i, row in enumerate(rows[1:], 1):
        st = row[status_col] if len(row) > status_col else ''
        c = color_map.get(st)
        if c:
            extra.append(('TEXTCOLOR', (status_col, i), (status_col, i), c))
            extra.append(('FONTNAME', (status_col, i), (status_col, i), PDF_FONT_BOLD))
    return TableStyle(base + extra)


# ─────────────────────────────────────────────────────────────────────────────
# EXCEL EXPORT
# ─────────────────────────────────────────────────────────────────────────────

def export_excel(report: dict, comparison_id: int = None) -> bytes:
    """Generuje raport Excel (.xlsx). Zwraca bytes."""
    if not HAS_OPENPYXL:
        raise ImportError("openpyxl nie jest zainstalowany. pip install openpyxl")

    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # usuń domyślny arkusz

    # ── Style ─────────────────────────────────────────────────────────────────
    def fill(hex_color):
        return PatternFill('solid', fgColor=hex_color.lstrip('#'))

    def font(bold=False, color='1a1d26', size=10, italic=False):
        return Font(name='Calibri', bold=bold, color=color, size=size, italic=italic)

    def border_thin():
        s = Side(style='thin', color='D0D4E0')
        return Border(left=s, right=s, top=s, bottom=s)

    def align(h='left', v='top', wrap=True):
        return Alignment(horizontal=h, vertical=v, wrap_text=wrap)

    HDR_FILL = fill('2a3050')
    HDR_FONT = font(bold=True, color='FFFFFF')
    ERR_FILL = fill('fde8e8')
    WARN_FILL= fill('fef3e2')
    FMT_FILL = fill('e8eeff')
    OK_FILL  = fill('e8faf0')
    ALT_FILL = fill('f8f9fc')

    def write_hdr(ws, row, cols):
        for c, text in enumerate(cols, 1):
            cell = ws.cell(row=row, column=c, value=text)
            cell.font = HDR_FONT
            cell.fill = HDR_FILL
            cell.alignment = align('center', 'center', False)
            cell.border = border_thin()

    def auto_width(ws, min_w=8, max_w=50):
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                try:
                    if cell.value:
                        max_len = max(max_len, len(str(cell.value)))
                except Exception:
                    pass
            ws.column_dimensions[col_letter].width = max(min_w, min(max_w, max_len + 2))

    now = datetime.now().strftime('%d.%m.%Y %H:%M')

    # ── ARKUSZ 1: Podsumowanie ────────────────────────────────────────────────
    ws = wb.create_sheet('Podsumowanie')
    ws.sheet_view.showGridLines = False

    # Nagłówek
    ws.merge_cells('A1:F1')
    ws['A1'] = 'DocCompare — Raport porównania'
    ws['A1'].font = Font(name='Calibri', bold=True, size=14, color='4F8EF7')
    ws['A1'].fill = fill('1e2130')
    ws['A1'].alignment = align('left', 'center', False)
    ws.row_dimensions[1].height = 28

    ws.merge_cells('A2:F2')
    ws['A2'] = f'ACME  |  {now}  |  Porównanie #{comparison_id or "—"}'
    ws['A2'].font = font(color='6B7080', italic=True)
    ws['A2'].fill = fill('f0f2f8')
    ws.row_dimensions[2].height = 16

    # Dane meta
    meta = [
        ('', ''),
        ('Dokument A:', report.get('file_a', '—')),
        ('Typ A:', report.get('doc_type_a', 'auto')),
        ('Dokument B:', report.get('file_b', '—')),
        ('Typ B:', report.get('doc_type_b', 'auto')),
        ('', ''),
        ('Status:', {'ok':'✅ ZGODNE','ostrzezenie':'⚠ OSTRZEŻENIA','blad':'❌ BŁĘDY'}
                    .get(report.get('risk_level','ok'),'?')),
        ('Rozbieżności:', report.get('total_errors', 0)),
        ('Ostrzeżenia:', report.get('total_warnings', 0)),
    ]
    for i, (k, v) in enumerate(meta, 3):
        ws.cell(row=i, column=1, value=k).font = font(bold=True, color='6B7080')
        c = ws.cell(row=i, column=2, value=v)
        c.font = font(bold=(k == 'Status:'))
        if k == 'Status:':
            risk = report.get('risk_level', 'ok')
            c.fill = {'ok': fill('e8faf0'), 'ostrzezenie': fill('fef3e2'),
                      'blad': fill('fde8e8')}.get(risk, ALT_FILL)

    auto_width(ws)
    ws.column_dimensions['A'].width = 18
    ws.column_dimensions['B'].width = 60

    # ── ARKUSZ 2: Pozycje towarowe ────────────────────────────────────────────
    tbl_m = report.get('modules', {}).get('table', {})
    if tbl_m.get('ok') and tbl_m.get('items'):
        ws2 = wb.create_sheet('Pozycje towarowe')
        ws2.sheet_view.showGridLines = False
        cols = ['REF', 'Opis A', 'Opis B', 'Ilość A', 'Ilość B',
                'Cena A', 'Cena B', 'Net A', 'Net B', 'Status', 'Problemy']
        write_hdr(ws2, 1, cols)
        ws2.freeze_panes = 'A2'
        ws2.auto_filter.ref = f'A1:{get_column_letter(len(cols))}1'

        status_fills = {
            'roznica': ERR_FILL, 'brak_w_a': ERR_FILL, 'brak_w_b': ERR_FILL,
            'format': FMT_FILL, 'watpliwe': WARN_FILL, 'ok': OK_FILL,
        }
        _all_items = tbl_m['items']
        _items_capped = _all_items[:MAX_ITEMS_PER_EXCEL]
        for r, it in enumerate(_items_capped, 2):
            st = it.get('status', '')
            row_fill = status_fills.get(st, ALT_FILL if r % 2 == 0 else None)
            issues_list = it.get('issues') or []
            vals = [
                _xsafe(it.get('ref')),
                _xsafe(it.get('desc_a')),
                _xsafe(it.get('desc_b')),
                _xsafe(it.get('qty_a')),
                _xsafe(it.get('qty_b')),
                _xsafe(it.get('price_a')),
                _xsafe(it.get('price_b')),
                _xsafe(it.get('net_a')),
                _xsafe(it.get('net_b')),
                st,
                '; '.join(str(x) for x in issues_list)[:120]
            ]
            for c, v in enumerate(vals, 1):
                cell = ws2.cell(row=r, column=c, value=v)
                cell.font = font()
                cell.border = border_thin()
                cell.alignment = align()
                if row_fill:
                    cell.fill = row_fill
        if len(_all_items) > MAX_ITEMS_PER_EXCEL:
            ws2.cell(row=len(_items_capped) + 2, column=1,
                     value=f"… ({len(_all_items) - MAX_ITEMS_PER_EXCEL} pozycji pominięto — "
                           f"limit {MAX_ITEMS_PER_EXCEL})").font = font()
        auto_width(ws2)

    # ── ARKUSZ 3: Pola nagłówkowe ─────────────────────────────────────────────
    if tbl_m.get('ok') and tbl_m.get('headers'):
        ws3 = wb.create_sheet('Pola nagłówkowe')
        ws3.sheet_view.showGridLines = False
        cols = ['Pole', 'Wartość A', 'Wartość B', 'Status', 'Komentarz']
        write_hdr(ws3, 1, cols)
        ws3.freeze_panes = 'A2'
        status_fills = {
            'roznica': ERR_FILL, 'brak_w_a': ERR_FILL, 'brak_w_b': ERR_FILL,
            'format': FMT_FILL, 'watpliwe': WARN_FILL, 'ok': OK_FILL,
        }
        for r, h in enumerate(tbl_m['headers'], 2):
            st = h.get('status','')
            row_fill = status_fills.get(st, ALT_FILL if r % 2 == 0 else None)
            vals = [_xsafe(h.get('key')), _xsafe(h.get('val_a')), _xsafe(h.get('val_b')),
                    _safe(st), _xsafe(h.get('comment'))]
            for c, v in enumerate(vals, 1):
                cell = ws3.cell(row=r, column=c, value=v)
                cell.font = font()
                cell.border = border_thin()
                cell.alignment = align()
                if row_fill:
                    cell.fill = row_fill
        auto_width(ws3)

    # ── ARKUSZ 4: Literówki ───────────────────────────────────────────────────
    typo_m = report.get('modules', {}).get('typo', {})
    if typo_m.get('ok') and typo_m.get('findings'):
        ws4 = wb.create_sheet('Literówki I↔1')
        ws4.sheet_view.showGridLines = False
        cols = ['Ważność', 'Kategoria', 'Pole', 'Wartość A', 'Wartość B', 'Opis', 'Sugestia']
        write_hdr(ws4, 1, cols)
        ws4.freeze_panes = 'A2'
        sev_fills = {'error': ERR_FILL, 'warning': WARN_FILL, 'info': FMT_FILL}
        for r, f in enumerate(typo_m['findings'], 2):
            sev = f.get('severity','')
            row_fill = sev_fills.get(sev, ALT_FILL if r % 2 == 0 else None)
            vals = [_safe(sev), _safe(f.get('category')), _xsafe(f.get('field_name')),
                    _xsafe(f.get('val_a')), _xsafe(f.get('val_b')),
                    _xsafe(f.get('description')), _xsafe(f.get('suggestion'))]
            for c, v in enumerate(vals, 1):
                cell = ws4.cell(row=r, column=c, value=v)
                cell.font = font()
                cell.border = border_thin()
                cell.alignment = align()
                if row_fill:
                    cell.fill = row_fill
        auto_width(ws4)

    # ── ARKUSZ 5: Pola tekstowe (regex) ──────────────────────────────────────
    rx_m = report.get('modules', {}).get('regex', {})
    if rx_m.get('ok') and rx_m.get('fields'):
        ws5 = wb.create_sheet('Pola tekstowe')
        ws5.sheet_view.showGridLines = False
        cols = ['Pole', 'Wartość A', 'Wartość B', 'Status', 'Komentarz']
        write_hdr(ws5, 1, cols)
        ws5.freeze_panes = 'A2'
        status_fills = {'roznica': ERR_FILL, 'brak_w_a': ERR_FILL, 'brak_w_b': ERR_FILL,
                        'format': FMT_FILL, 'watpliwe': WARN_FILL, 'ok': OK_FILL}
        for r, f in enumerate(rx_m['fields'], 2):
            st = f.get('status','')
            row_fill = status_fills.get(st, ALT_FILL if r % 2 == 0 else None)
            vals = [_xsafe(f.get('field')), _xsafe(f.get('doc_a')), _xsafe(f.get('doc_b')),
                    _safe(st), _xsafe(f.get('comment'))]
            for c, v in enumerate(vals, 1):
                cell = ws5.cell(row=r, column=c, value=v)
                cell.font = font()
                cell.border = border_thin()
                cell.alignment = align()
                if row_fill:
                    cell.fill = row_fill
        auto_width(ws5)

    buf = io.BytesIO()
    wb.save(buf)
    wb.close()
    return buf.getvalue()
