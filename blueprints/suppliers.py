"""
blueprints/suppliers.py — dostawcy: CRUD, wizard profili, master data (Phase 3, plan 03-02).

Wydzielone z app.py (trzy nieciągłe pasma linii — zob. 03-RESEARCH.md Pitfall 5).
Importuje wyłącznie core/db + moduły logiki biznesowej już wydzielone
(supplier_master.py, supplier_profiles.py) — blueprint → core/db, nigdy → app.
"""
from __future__ import annotations

import json
import logging

from flask import Blueprint, render_template, request, jsonify, session
from pydantic import BaseModel, ConfigDict, Field, ValidationError, field_validator, model_validator

from db import get_db
from core.security import login_required, require_role, csrf_protect, is_admin
from core.audit import log_audit as _log_audit
from core.validation import validate_body
from pdf_validation import validate_pdf_upload as _validate_pdf_upload

logger = logging.getLogger("doccompare")

bp = Blueprint("suppliers", __name__)   # NO url_prefix — paths are heterogeneous


# ── Pydantic models (REFACTOR-03) ──────────────────────────────────────────────
# Mass-assignment guard: none of these declare id/sid/code — sid/code are always
# taken from the URL path, never the body (see 03-02-PLAN.md prohibitions).

class SupplierCreate(BaseModel):
    """POST /api/suppliers — only gates the one field this route itself requires
    (`name`); `code` stays a manual check below (mass-assignment guard forbids a
    `code` field here) and the full raw body is still forwarded to
    supplier_profiles.save_supplier(), which already tolerates arbitrary keys."""
    name: str = Field(min_length=1, max_length=200)


class SupplierUpdate(BaseModel):
    """PATCH /api/suppliers/<sid> — preserves the original 'if not data: 400'
    gate (empty body rejected) while allowing the full set of wizard fields
    through untouched (extra="allow")."""
    model_config = ConfigDict(extra="allow")

    @model_validator(mode="before")
    @classmethod
    def _reject_empty(cls, data):
        if not data:
            raise ValueError("Brak danych JSON")
        return data


class SupplierDetect(BaseModel):
    """POST /api/suppliers/detect — text is optional (empty text is a valid,
    non-error request per the original behavior: returns {"supplier": None})."""
    text: str = ""


class SupplierRulesUpdate(BaseModel):
    """POST /api/suppliers/<sid>/rules — type + max-length(500) gate; per-item
    `type` enum + string truncation stay inline (unchanged from before)."""
    rules: list[dict] = Field(default_factory=list, max_length=500)


class SupplierContactUpdate(BaseModel):
    """PATCH /api/sup-master/<code>/contact — mirrors 03-RESEARCH.md's exemplar:
    manual '@'/'.' email check kept inside a plain str field (no email-validator
    dependency, per YAGNI)."""
    email: str | None = None
    contact_person: str | None = Field(default=None, max_length=200)
    phone: str | None = Field(default=None, max_length=50)

    @field_validator("email")
    @classmethod
    def _check_email(cls, v):
        if v is not None:
            v = str(v).strip()
            if v and ("@" not in v or "." not in v.split("@")[-1]):
                raise ValueError("Niepoprawny adres e-mail")
        return v


def _supplier_doc_config(mapping: dict, required=("ref", "qty", "price")) -> dict:
    """Analizuje mapowanie kolumn jednego doc-type. Zwraca {field_count, roles, ok, partial}."""
    if not mapping or not isinstance(mapping, dict):
        return {"field_count": 0, "roles": [], "ok": False, "partial": False}
    roles = set(v for v in mapping.values() if v and v != "skip")
    field_count = len(roles)
    has_all = all(r in roles for r in required)
    has_ref = "ref" in roles
    return {
        "field_count": field_count,
        "roles": sorted(roles),
        "ok": has_all,
        "partial": has_ref and not has_all,
    }


@bp.route("/suppliers")
@login_required
def suppliers_page():
    from supplier_profiles import get_all_suppliers
    import json as _json
    suppliers_raw = get_all_suppliers()

    db = get_db()
    try:
        comp_rows = db.execute(
            "SELECT supplier_code, COUNT(*) as cnt FROM comparisons "
            "WHERE supplier_code IS NOT NULL AND supplier_code != '' "
            "GROUP BY supplier_code"
        ).fetchall()
    finally:
        db.close()
    comp_counts = {r["supplier_code"]: r["cnt"] for r in comp_rows}

    suppliers = []
    for s in suppliers_raw:
        s2 = dict(s)
        # Deserializuj pola JSON jeśli nadal string
        for fld in ["pi_column_mapping_json", "po_column_mapping_json",
                    "sad_column_mapping_json",
                    "detect_keywords_json", "payment_terms_json",
                    "ignore_fields_json", "custom_rules_json",
                    "column_mapping_json", "known_issues_json"]:
            v = s2.get(fld)
            if isinstance(v, str) and v:
                try:
                    s2[fld] = _json.loads(v)
                except Exception:
                    s2[fld] = {} if "{" in (v or "") else []

        # Fallback: wyciągnij PI/PO mapping z column_mapping_json dla starych profili
        col_map = s2.get("column_mapping_json") or {}
        for dtype, key in [("PI","pi_column_mapping_json"),("PO","po_column_mapping_json"),
                           ("SAD","sad_column_mapping_json"),("CI","ci_column_mapping_json"),
                           ("PL","pl_column_mapping_json")]:
            if not s2.get(key) and isinstance(col_map, dict):
                s2[key] = col_map.get(dtype, col_map.get(dtype.lower(), {}))

        kw = s2.get("detect_keywords_json")
        if not isinstance(kw, list):
            s2["detect_keywords_json"] = [kw] if isinstance(kw, str) and kw else []
        ki = s2.get("known_issues_json")
        if not isinstance(ki, list):
            s2["known_issues_json"] = []

        # Per-doc-type config analysis
        s2["_doc_config"] = {
            "PI":  _supplier_doc_config(s2.get("pi_column_mapping_json") or {}),
            "PO":  _supplier_doc_config(s2.get("po_column_mapping_json") or {}),
            "SAD": _supplier_doc_config(s2.get("sad_column_mapping_json") or {},
                                        required=("ref", "tariff_code")),
            "CI":  _supplier_doc_config(s2.get("ci_column_mapping_json") or {}),
            "PL":  _supplier_doc_config(s2.get("pl_column_mapping_json") or {}),
        }
        # Overall config status
        pi_ok = s2["_doc_config"]["PI"]["ok"]
        po_ok = s2["_doc_config"]["PO"]["ok"]
        any_mapped = any(v["field_count"] > 0 for v in s2["_doc_config"].values())
        s2["_config_status"] = "full" if (pi_ok and po_ok) else ("partial" if any_mapped else "none")

        # Comparison count
        s2["comparison_count"] = comp_counts.get(s2.get("code", ""), 0)

        suppliers.append(s2)

    return render_template("suppliers.html",
                           suppliers=suppliers,
                           username=session["username"],
                           role=session["role"],
                           is_admin=is_admin(session["role"]))


@bp.route("/api/suppliers")
@login_required
def api_suppliers_list():
    from supplier_profiles import get_all_suppliers
    return jsonify(get_all_suppliers())


@bp.route("/api/suppliers", methods=["POST"])
@require_role("admin")
@csrf_protect
@validate_body(SupplierCreate)
def api_suppliers_add(body: SupplierCreate):
    from supplier_profiles import save_supplier
    data = request.get_json(silent=True) or {}
    if not data.get("code"):
        return jsonify({"error": "Wymagane: code, name"}), 400
    try:
        result = save_supplier(data)
        return jsonify({"ok": True, "supplier": result})
    except Exception as e:
        logger.exception("save_supplier error: %s", e)
        return jsonify({"error": "Błąd zapisu dostawcy — sprawdź dane."}), 400


@bp.route("/api/suppliers/<int:sid>", methods=["PATCH"])
@require_role("admin")
@csrf_protect
def api_suppliers_update(sid):
    # WR-01: row lookup (404) runs before body validation (400) — restores the
    # pre-refactor priority (@validate_body would run before this function body,
    # flipping the order for a nonexistent id + empty body).
    from supplier_profiles import get_all_suppliers, save_supplier
    db = get_db()
    try:
        row = db.execute("SELECT * FROM suppliers WHERE id=?", (sid,)).fetchone()
    finally:
        db.close()
    if not row:
        return jsonify({"error": "Nie znaleziono"}), 404
    try:
        SupplierUpdate.model_validate(request.get_json(silent=True) or {})
    except ValidationError as e:
        first = e.errors()[0]
        field = ".".join(str(p) for p in first["loc"])
        return jsonify({"error": f"Pole '{field}': {first['msg']}"}), 400
    data = request.get_json(silent=True)
    data["code"] = dict(row)["code"]  # zachowaj oryginalny kod
    try:
        result = save_supplier(data)
        return jsonify({"ok": True, "supplier": result})
    except Exception as e:
        logger.exception("update_supplier error: %s", e)
        return jsonify({"error": "Błąd aktualizacji dostawcy — sprawdź dane."}), 400


@bp.route("/api/suppliers/<int:sid>", methods=["DELETE"])
@require_role("admin")
@csrf_protect
def api_suppliers_delete(sid):
    from supplier_profiles import delete_supplier
    ok = delete_supplier(sid)
    if not ok:
        return jsonify({"error": "Nie można usunąć wbudowanego profilu"}), 400
    return jsonify({"ok": True})


@bp.route("/api/suppliers/detect", methods=["POST"])
@login_required
@csrf_protect
@validate_body(SupplierDetect)
def api_suppliers_detect(body: SupplierDetect):
    """Wykrywa dostawcę z tekstu PDF."""
    if not body.text:
        return jsonify({"supplier": None})
    from supplier_profiles import detect_supplier
    s = detect_supplier(body.text)
    return jsonify({"supplier": s})


# ─────────────────────────────────────────────────────────────────────────────
# WIZARD PROFILI DOSTAWCÓW
# ─────────────────────────────────────────────────────────────────────────────

@bp.route("/suppliers/wizard/new")
@require_role("admin")
def supplier_wizard_new():
    return render_template("supplier_wizard.html",
                           mode="new", supplier=None, supplier_id=None,
                           supplier_name="",
                           username=session["username"],
                           role=session["role"])


@bp.route("/suppliers/wizard/<int:sid>")
@require_role("admin")
def supplier_wizard_edit(sid):
    db = get_db()
    try:
        row = db.execute("SELECT * FROM suppliers WHERE id=?", (sid,)).fetchone()
    finally:
        db.close()
    if not row:
        return "Nie znaleziono dostawcy", 404
    s = dict(row)
    # Deserializuj JSON pola
    import json as _json
    for fld in ["pi_column_mapping_json", "po_column_mapping_json",
                "sad_column_mapping_json",
                "detect_keywords_json", "payment_terms_json",
                "ignore_fields_json", "custom_rules_json"]:
        if s.get(fld) and isinstance(s[fld], str):
            try:
                s[fld] = _json.loads(s[fld])
            except Exception:
                pass
    # Zbuduj payment_terms_text
    pt = s.get("payment_terms_json") or {}
    if isinstance(pt, dict):
        s["payment_terms_text"] = "\n".join(f"{k}={v}" for k,v in pt.items())
    else:
        s["payment_terms_text"] = ""
    # ignore_fields jako string
    ig = s.get("ignore_fields_json") or []
    s["ignore_fields"] = ", ".join(ig) if isinstance(ig, list) else str(ig)
    return render_template("supplier_wizard.html",
                           mode="edit", supplier=s,
                           supplier_id=sid,
                           supplier_name=s.get("name",""),
                           username=session["username"],
                           role=session["role"])


@bp.route("/api/suppliers/preview_table", methods=["POST"])
@require_role("admin")
@csrf_protect
def api_suppliers_preview_table():
    """
    Przyjmuje plik PDF, zwraca podgląd tabeli produktów z auto-detekcją kolumn.
    Używany przez wizard mapowania.
    """
    import traceback
    import tempfile, os
    file = request.files.get("file")
    side = request.form.get("side", "pi")  # 'pi' lub 'po'
    if not file:
        return jsonify({"error": "Brak pliku"}), 400
    err = _validate_pdf_upload(file)
    if err:
        return jsonify({"error": err}), 400
    # Zapisz tymczasowo
    suffix = ".pdf"
    with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as tmp:
        file.save(tmp.name)
        tmp_path = tmp.name
    try:
        from supplier_profiles import extract_table_preview
        result = extract_table_preview(tmp_path, doc_type=side.upper())
        # Dodaj przykłady REF jeśli znaleziono kolumnę ref
        if result.get("table_found") and result.get("rows"):
            ref_col = None
            dm = result.get("detected_mappings", {})
            # Znajdź indeks kolumny ref
            for hdr, role in dm.items():
                if role == "ref":
                    try:
                        ref_col = result["headers"].index(hdr)
                    except ValueError:
                        pass
                    break
            if ref_col is not None:
                samples = []
                for row in result["rows"][:4]:
                    if ref_col < len(row) and row[ref_col]:
                        samples.append(str(row[ref_col])[:30])
                result["ref_samples"] = samples
        return jsonify(result)
    except Exception as e:
        tb = traceback.format_exc()
        logger.error("preview_table error: %s", tb)
        return jsonify({"error": "Błąd podglądu tabeli.", "table_found": False}), 500
    finally:
        try:
            os.unlink(tmp_path)
        except Exception:
            pass


@bp.route("/api/suppliers/<int:sid>/audit")
@require_role("admin")
def api_suppliers_audit(sid):
    """Historia zmian profilu dostawcy."""
    from supplier_profiles import get_audit_log
    log = get_audit_log(sid)
    # Dodaj pole 'summary' dla template
    result = []
    for entry in log:
        diff = entry.get('diff_json', {})
        if isinstance(diff, dict):
            keys = list(diff.keys())[:3]
            summary = ', '.join(keys) + (' …' if len(diff) > 3 else '')
        else:
            summary = str(diff)[:80]
        result.append({
            'changed_at': entry.get('changed_at',''),
            'changed_by': entry.get('changed_by','system'),
            'action':     entry.get('action','update'),
            'summary':    summary,
            'changes_json': diff,
        })
    return jsonify(result)


@bp.route("/api/suppliers/<int:sid>/stats")
@login_required
def api_suppliers_stats(sid):
    """Statystyki użycia profilu — ile porównań, % błędów."""
    db = get_db()
    try:
        row = db.execute("SELECT * FROM suppliers WHERE id=?", (sid,)).fetchone()
        if not row:
            return jsonify({"error": "Nie znaleziono"}), 404
        supplier_code = dict(row).get("code", "")
        try:
            stats = db.execute(
                """SELECT COUNT(*) as total,
                          SUM(CASE WHEN diff_count > 0 THEN 1 ELSE 0 END) as with_errors,
                          MAX(created_at) as last_used
                   FROM comparisons
                   WHERE supplier_code=? AND doc_type != 'artwork'""",
                (supplier_code,)
            ).fetchone()
            total = stats["total"] or 0
            errors = stats["with_errors"] or 0
            return jsonify({
                "total_comparisons": total,
                "error_rate_pct": round(errors / total * 100, 1) if total else 0,
                "last_used": stats["last_used"],
            })
        except Exception:
            return jsonify({"total_comparisons": 0, "error_rate_pct": 0, "last_used": None})
    finally:
        db.close()


# ═══════════════════════════════════════════════════════════════
# FUNKCJA 16 — REGUŁY BIZNESOWE PER DOSTAWCA
# ═══════════════════════════════════════════════════════════════

@bp.route("/api/suppliers/<int:sid>/rules", methods=["GET"])
@login_required
def get_supplier_rules(sid):
    db = get_db()
    try:
        row = db.execute("SELECT custom_rules_json FROM suppliers WHERE id=?", (sid,)).fetchone()
    finally:
        db.close()
    if not row:
        return jsonify({"error": "Nie znaleziono"}), 404
    try:
        rules = json.loads(row["custom_rules_json"] or "[]")
    except (ValueError, TypeError):
        rules = []
    return jsonify({"rules": rules})


@bp.route("/api/suppliers/<int:sid>/rules", methods=["POST"])
@require_role("admin")
@csrf_protect
@validate_body(SupplierRulesUpdate)
def save_supplier_rules(body: SupplierRulesUpdate, sid):
    rules = body.rules
    # Waliduj typy
    valid_types = {"PRICE_ROUNDING","CURRENCY_TOLERANCE","FIELD_IGNORE",
                   "PAYMENT_ALIAS","REF_TRANSFORM","QTY_MULTIPLIER"}
    _rule_str_cap = 500
    sanitized_rules = []
    for r in rules:
        if not isinstance(r, dict) or r.get("type") not in valid_types:
            return jsonify({"error": f"Nieznany typ reguły: {r.get('type') if isinstance(r, dict) else r}"}), 400
        sanitized_rules.append({k: str(v)[:_rule_str_cap] if isinstance(v, str) else v
                                 for k, v in r.items()})
    rules = sanitized_rules
    db = get_db()
    try:
        db.execute("UPDATE suppliers SET custom_rules_json=? WHERE id=?",
                   (json.dumps(rules), sid))
        db.commit()
    finally:
        db.close()
    return jsonify({"ok": True, "rules_count": len(rules)})


# ── Master data dostawców (rozszerza tabelę suppliers) ─────────────────────────
@bp.route("/suppliers/master")
@login_required
def page_suppliers_master():
    return render_template("suppliers_master.html",
                           username=session.get("username"), role=session.get("role"))


@bp.route("/api/sup-master", methods=["GET"])
@login_required
def api_sup_master_list():
    import supplier_master as _sm
    db = get_db()
    try:
        items = _sm.list_suppliers(db, request.args.get("q", "").strip())
    finally:
        db.close()
    return jsonify({"items": items, "total": len(items)})


@bp.route("/api/sup-master/<path:code>", methods=["GET"])
@login_required
def api_sup_master_detail(code):
    import supplier_master as _sm
    db = get_db()
    try:
        s = _sm.get_supplier(db, code)
    finally:
        db.close()
    if not s:
        return jsonify({"error": "Nie znaleziono dostawcy"}), 404
    try:
        kw = json.loads(s.get("detect_keywords_json") or "[]")
    except (ValueError, TypeError):
        kw = []
    return jsonify({
        "code": s.get("code"), "name": s.get("name"),
        "producer_code": s.get("producer_code"), "country": s.get("country"),
        "currency": s.get("currency"), "incoterms": s.get("incoterms"),
        "payment_terms_default": s.get("payment_terms_default"),
        "lead_time_days": s.get("lead_time_days"),
        "contact_person": s.get("contact_person"), "email": s.get("email"),
        "phone": s.get("phone"), "notes": s.get("notes"), "keywords": kw,
    })


@bp.route("/api/sup-master/<path:code>/contact", methods=["PATCH"])
@require_role("manager")
@csrf_protect
@validate_body(SupplierContactUpdate)
def api_sup_master_update_contact(body: SupplierContactUpdate, code):
    """Ustaw/edytuj dane kontaktowe dostawcy (e-mail odbiorcy zamówień, osoba, tel.)."""
    import supplier_master as _sm
    db = get_db()
    try:
        res = _sm.update_contact(
            db, code, email=body.email,
            contact_person=body.contact_person, phone=body.phone)
    finally:
        db.close()
    if "error" in res:
        return jsonify(res), 400
    _log_audit("supplier_contact_update", session.get("username"),
               f"{code} email={body.email}")
    return jsonify(res)


@bp.route("/api/sup-master/import", methods=["POST"])
@require_role("manager")
@csrf_protect
def api_sup_master_import():
    import supplier_master as _sm
    f = request.files.get("file")
    if not f or not f.filename:
        return jsonify({"error": "Brak pliku"}), 400
    fname = f.filename.lower()
    raw_rows = []
    try:
        if fname.endswith(".xlsx"):
            from openpyxl import load_workbook
            wb = load_workbook(f, read_only=True, data_only=True)
            ws = wb.active
            for r in ws.iter_rows(values_only=True):
                raw_rows.append(list(r))
                if len(raw_rows) > 20000:
                    break
        elif fname.endswith(".csv"):
            import csv as _csv, io as _io
            raw = f.read().decode("utf-8-sig", errors="replace")
            # Wykryj separator (polski Excel zapisuje CSV ze średnikiem).
            _sample = raw[:8192]
            _delim = max((",", ";", "\t"), key=lambda d: _sample.count(d))
            raw_rows = [r for r in _csv.reader(_io.StringIO(raw), delimiter=_delim)]
        else:
            return jsonify({"error": "Obsługiwane formaty: .xlsx, .csv"}), 400
    except Exception as _e:
        logger.warning("supplier master import parse error: %s", _e)
        return jsonify({"error": "Nie udało się odczytać pliku"}), 400
    db = get_db()
    try:
        res = _sm.import_workbook(db, raw_rows)
    finally:
        db.close()
    _log_audit("sup_master_import", session.get("username"), f"imported={res['imported']}")
    return jsonify({"ok": True, **res})
