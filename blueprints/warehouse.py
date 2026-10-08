"""
blueprints/warehouse.py — magazyn: stany magazynowe (lista/import) i rejestr
przyjęć (Phase 3, plan 03-03).

Wydzielone z app.py (dwa nieciągłe pasma linii — zob. 03-RESEARCH.md Pitfall 5).
Logika biznesowa już wydzielona do warehouse_stock.py — trasy to cienkie
opakowanie (blueprint → core/db, nigdy → app).
"""
from __future__ import annotations

import logging
from typing import Literal

from flask import Blueprint, render_template, request, jsonify, session
from pydantic import BaseModel, Field

from db import get_db
from core.security import login_required, require_role, csrf_protect, can_see_all
from core.audit import log_audit as _log_audit
from core.validation import validate_body
from delivery_workflow import PHASES as _DELIVERY_PHASES
import warehouse_stock as _ws

logger = logging.getLogger("doccompare")

bp = Blueprint("warehouse", __name__)   # NO url_prefix — paths are heterogeneous


# ── Pydantic models (REFACTOR-03) ──────────────────────────────────────────────
# Mass-assignment guard: no rid/id field — rid always comes from the URL path.

class WarehouseReceiptUpdate(BaseModel):
    """PUT /api/warehouse/receipts/<rid> — all fields Optional so partial
    updates work via model_dump(exclude_unset=True); status is a closed enum
    and unload_norm_min must be >= 0 (the pre-Pydantic code silently clamped
    negative values to 0 instead of rejecting them)."""
    container_number: str | None = None
    po_numbers: str | None = None
    received_date: str | None = None
    carrier: str | None = None
    unload_location: str | None = None
    condition_notes: str | None = None
    sap_document: str | None = None
    notes: str | None = None
    unloaded_by: str | None = None
    unload_start: str | None = None
    unload_end: str | None = None
    unload_norm_min: int | None = Field(default=None, ge=0)
    status: Literal["przybyle", "sprawdzone", "do_sap", "wprowadzone_sap"] | None = None


@bp.route("/warehouse/stock")
@login_required
def page_warehouse_stock():
    return render_template("warehouse_stock.html",
                           username=session.get("username"), role=session.get("role"))


@bp.route("/api/warehouse/stock", methods=["GET"])
@login_required
def api_warehouse_stock_list():
    db = get_db()
    try:
        return jsonify({"items": _ws.list_stock(db)})
    finally:
        db.close()


@bp.route("/api/warehouse/stock/import", methods=["POST"])
@require_role("manager")
@csrf_protect
def api_warehouse_stock_import():
    """Import stanów magazynowych z CSV/XLSX (REF + stan + zużycie miesięczne).
    Upsert po znormalizowanym REF — zasila model „na ile wystarczy" w analityce."""
    f = request.files.get("file")
    if not f or not f.filename:
        return jsonify({"error": "Brak pliku"}), 400
    fname = f.filename.lower()
    raw_rows = []
    try:
        if fname.endswith(".xlsx"):
            from openpyxl import load_workbook
            wb = load_workbook(f, read_only=True, data_only=True)
            ws = wb.active
            for r in ws.iter_rows(values_only=True):
                raw_rows.append(list(r))
                if len(raw_rows) > 60000:
                    break
        elif fname.endswith(".csv"):
            import csv as _csv, io as _io
            raw = f.read().decode("utf-8-sig", errors="replace")
            _sample = raw[:8192]
            _delim = max((",", ";", "\t"), key=lambda d: _sample.count(d))
            raw_rows = [r for r in _csv.reader(_io.StringIO(raw), delimiter=_delim)]
        else:
            return jsonify({"error": "Obsługiwane formaty: .xlsx, .csv"}), 400
    except Exception as _e:
        logger.warning("warehouse stock import parse error: %s", _e)
        return jsonify({"error": "Nie udało się odczytać pliku"}), 400
    if len(raw_rows) > 55000:
        return jsonify({"error": "Za dużo wierszy (max ~55000)"}), 400
    db = get_db()
    try:
        res = _ws.import_rows(db, raw_rows)
    finally:
        db.close()
    if res.get("error"):
        return jsonify(res), 400
    _log_audit("warehouse_stock_import", session.get("username"),
               f"imported={res.get('imported')} skipped={res.get('skipped')}")
    return jsonify({"ok": True, **res})


def _parse_date_only(s):
    """Parsuje datę (różne formaty) do date albo None — do okna 'nadchodzące'."""
    from datetime import datetime as _dt
    s = str(s or "").strip()
    if not s:
        return None
    s = s.replace("T", " ").split(" ")[0]
    for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%d-%m-%Y", "%Y/%m/%d", "%d/%m/%Y"):
        try:
            return _dt.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


@bp.route("/api/warehouse/upcoming")
@login_required
def api_warehouse_upcoming():
    """Nadchodzące dostawy w bliskim oknie (domyślnie 14 dni) wg najbliższej daty
    dostawy / ETA / ETD — bez tych już odebranych (faza magazyn i dalej). Per dział."""
    from datetime import date, timedelta
    try:
        days = int(request.args.get("days", 14))
    except (TypeError, ValueError):
        days = 14
    days = max(1, min(days, 120))
    db = get_db()
    try:
        dept_clause, dept_params = "", []
        if not can_see_all(session.get("role", "")):
            dept_clause = " AND (COALESCE(z.department,'')='' OR z.department=?)"
            dept_params = [session.get("department") or ""]
        _mag_idx = next((i for i, ph in enumerate(_DELIVERY_PHASES)
                         if "magaz" in ph["label"].lower()), len(_DELIVERY_PHASES))
        received = set()
        for i, ph in enumerate(_DELIVERY_PHASES):
            if i >= _mag_idx:
                received.update(ph["statuses"])
        rows = db.execute(
            # bandit: WHERE/SQL składany ze stałych fragmentów w kodzie; wartości przez ?
            "SELECT z.nr_zamowienia AS nr, z.supplier_name AS sup, z.status AS st, "  # nosec B608
            "z.data_dostawy AS dd, z.planowane_etd AS etd, z.department AS dept, "
            "k.numer_kontenera AS kont, k.eta AS eta "
            "FROM kolejka_zlecenia z LEFT JOIN kolejka_kontenery k ON z.kontener_id=k.id "
            "WHERE 1=1" + dept_clause, dept_params).fetchall()
    finally:
        db.close()
    today = date.today()
    horizon = today + timedelta(days=days)
    past = today - timedelta(days=3)  # lekko spóźnione też pokazujemy
    out = []
    for r in rows:
        if (r["st"] or "") in received:
            continue
        d_dd, d_eta, d_etd = (_parse_date_only(r["dd"]), _parse_date_only(r["eta"]),
                              _parse_date_only(r["etd"]))
        d = d_dd or d_eta or d_etd
        if not d or d < past or d > horizon:
            continue
        out.append({
            "nr": r["nr"], "supplier": r["sup"], "status": r["st"],
            "container": r["kont"] or "", "arrival": d.isoformat(),
            "days_left": (d - today).days,
            "source": ("data dostawy" if d_dd else ("ETA" if d_eta else "ETD")),
            "department": r["dept"] or "",
        })
    out.sort(key=lambda x: x["arrival"])
    return jsonify({"items": out, "days": days, "count": len(out)})


@bp.route("/warehouse")
@login_required
def warehouse_page():
    db = get_db()
    try:
        rows = db.execute(
            "SELECT wr.*, u.username as created_by_name, u2.username as checked_by_name "
            "FROM warehouse_receipts wr "
            "LEFT JOIN users u ON wr.created_by=u.id "
            "LEFT JOIN users u2 ON wr.checked_by=u2.id "
            "ORDER BY wr.created_at DESC LIMIT 200"
        ).fetchall()
        items = [dict(r) for r in rows]
    finally:
        db.close()
    return render_template("warehouse.html",
                           items=items,
                           username=session.get("username"),
                           role=session.get("role"))


@bp.route("/api/warehouse/receipts", methods=["POST"])
@require_role("manager")
@csrf_protect
def api_warehouse_create():
    # Ręczne dodawanie przyjęć jest wyłączone — wpisy powstają automatycznie, gdy
    # dostawa wejdzie w status magazynowy (_ensure_warehouse_receipt_for_delivery).
    # Operator wyłącznie edytuje istniejące wpisy (PUT /api/warehouse/receipts/<id>).
    return jsonify({
        "error": "Przyjęcia magazynowe tworzone są automatycznie, gdy dostawa "
                 "wejdzie w status magazynowy. Edytuj istniejący wpis zamiast "
                 "dodawać nowy."
    }), 405


@bp.route("/api/warehouse/receipts/<int:rid>", methods=["PUT"])
@require_role("manager")
@csrf_protect
@validate_body(WarehouseReceiptUpdate)
def api_warehouse_update(body: WarehouseReceiptUpdate, rid):
    data = body.model_dump(exclude_unset=True)
    uid = session["user_id"]
    _wh_max_lens = {
        "container_number": 50, "po_numbers": 500, "received_date": 20,
        "carrier": 200, "unload_location": 200, "status": 50,
        "condition_notes": 1000, "sap_document": 50, "notes": 1000,
        "unloaded_by": 200, "unload_start": 30, "unload_end": 30,
    }
    _wh_int_fields = {"unload_norm_min"}
    allowed = ["container_number", "po_numbers", "received_date", "carrier",
               "unload_location", "status", "condition_notes", "sap_document", "notes",
               "unloaded_by", "unload_start", "unload_end", "unload_norm_min"]
    fields, vals = [], []
    for k in allowed:
        if k in data and data[k] is not None:
            # explicit JSON null means "not provided" here too — `exclude_unset`
            # only drops absent keys, not `{"status": null}` (CR-01)
            v = data[k]
            if k in _wh_int_fields:
                v = int(v)  # already validated ge=0 by WarehouseReceiptUpdate
            else:
                v = str(v or "").strip()[:_wh_max_lens.get(k, 500)]
            fields.append(f"{k}=?")
            vals.append(v)
    if not fields:
        return jsonify({"error": "Brak pól"}), 400
    # If status changes to 'sprawdzone', record checker
    if data.get("status") == "sprawdzone":
        fields.append("checked_by=?")
        vals.append(uid)
        fields.append("checked_at=datetime('now')")
    fields.append("updated_at=datetime('now')")
    vals.append(rid)
    db = get_db()
    try:
        # bandit: nazwy kolumn z białej listy w kodzie (nie z kluczy żądania); wartości przez ?
        db.execute(f"UPDATE warehouse_receipts SET {', '.join(fields)} WHERE id=?", vals)  # nosec B608
        db.commit()
        # On wprowadzone_sap → mark matching shipments as zakonczone
        if data.get("status") == "wprowadzone_sap":
            try:
                row = db.execute(
                    "SELECT po_numbers FROM warehouse_receipts WHERE id=?", (rid,)
                ).fetchone()
                if row:
                    for _po in [p.strip() for p in (row["po_numbers"] or "").split(",") if p.strip()]:
                        db.execute(
                            "UPDATE shipments SET status='zakonczone', updated_at=datetime('now') "
                            "WHERE po_number=? AND status != 'zakonczone'",
                            (_po,)
                        )
                    db.commit()
            except Exception:
                pass
        return jsonify({"ok": True})
    finally:
        db.close()


@bp.route("/api/warehouse/receipts/<int:rid>", methods=["DELETE"])
@require_role("manager")
@csrf_protect
def api_warehouse_delete(rid):
    db = get_db()
    try:
        db.execute("DELETE FROM warehouse_receipts WHERE id=?", (rid,))
        db.commit()
        return jsonify({"ok": True})
    finally:
        db.close()
