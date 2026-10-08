#!/usr/bin/env python3
"""
tools/gen_master_data_template.py — generuje szablon Excel do importu master daty
materiałów (REF + opis PL/EN + EAN + rodzina + wymiary/EAN/artwork per poziom
opakowania: sztuka / OP / OPZ / karton).

Nagłówki są nazwami maszynowymi (pod import CSV/XLSX); komentarze w komórkach
nagłówka opisują znaczenie, a osobny arkusz „Instrukcja" tłumaczy format.

Uruchom:  python tools/gen_master_data_template.py [ścieżka_wyjściowa.xlsx]
Wymaga:   openpyxl
"""
import sys

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter

BASE = "DDEBF7"
LEVELS = [("sztuka", "Sztuka (JED)", "E2EFDA"), ("op", "Opakowanie (OP)", "DDEBF7"),
          ("opz", "OPZ", "E4DFEC"), ("karton", "Karton (KAR)", "FCE4D6")]


def build_columns():
    cols = [
        ("ref_code", "REF / indeks produktu (z zamówienia). WYMAGANE. Klucz unikalny.", 16, BASE),
        ("opis_pl",  "Opis produktu po polsku.", 34, BASE),
        ("opis_en",  "Opis produktu po angielsku.", 34, BASE),
        ("ean",      "Główny EAN-13 produktu (opcjonalnie).", 16, BASE),
        ("rodzina",  "Rodzina / grupa produktowa (opcjonalnie).", 18, BASE),
    ]
    for code, label, color in LEVELS:
        cols += [
            (f"{code}_wymiar",      f"[{label}] Wymiary: dł x szer x wys (mm). Np. 150 x 80 x 40", 18, color),
            (f"{code}_ean",         f"[{label}] EAN-13 tego poziomu (jeśli inny niż główny).", 16, color),
            (f"{code}_artwork_ref", f"[{label}] REF artworku w systemie. Wpisz, gdy nazwa pliku artworku "
                                    f"różni się od ref_code. Puste = brak mapowania.", 20, color),
        ]
    return cols


def generate(out_path: str):
    wb = Workbook()
    ws = wb.active
    ws.title = "Dane materiałowe"
    cols = build_columns()

    hdr_font = Font(bold=True, size=10, color="1F3864")
    thin = Side(style="thin", color="B0B0B0")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for i, (name, desc, width, color) in enumerate(cols, start=1):
        c = ws.cell(row=1, column=i, value=name)
        c.font = hdr_font
        c.fill = PatternFill("solid", fgColor=color)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = border
        cm = Comment(desc, "DocCompare"); cm.width = 260; cm.height = 110
        c.comment = cm
        ws.column_dimensions[get_column_letter(i)].width = width

    ex1 = {
        "ref_code": "RNBM10001", "opis_pl": "Rękawice nitrylowe bezpudrowe M",
        "opis_en": "Nitrile examination gloves, powder-free, M", "ean": "5900010800001", "rodzina": "Rękawice",
        "sztuka_wymiar": "240 x 120 x 5 mm", "sztuka_ean": "5900010800001", "sztuka_artwork_ref": "RNBM10001",
        "op_wymiar": "245 x 125 x 70 mm", "op_ean": "5900010800018", "op_artwork_ref": "RNBM10001",
        "karton_wymiar": "520 x 380 x 300 mm", "karton_ean": "5900010800025", "karton_artwork_ref": "RNBM10001-KARTON",
    }
    ex2 = {"ref_code": "AT-NF-S 15", "opis_pl": "Serweta niejałowa 15x15", "opis_en": "Non-sterile drape 15x15",
           "rodzina": "Serwety", "karton_wymiar": "400 x 300 x 250 mm", "karton_artwork_ref": "AT-NF-S 15"}
    name_to_idx = {n: i for i, (n, _, _, _) in enumerate(cols, start=1)}
    for r, ex in enumerate([ex1, ex2], start=2):
        for name, idx in name_to_idx.items():
            cell = ws.cell(row=r, column=idx, value=ex.get(name, ""))
            cell.border = border
            cell.font = Font(size=10, color="808080" if r == 3 else "000000")

    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 42

    ins = wb.create_sheet("Instrukcja")
    ins.column_dimensions["A"].width = 26
    ins.column_dimensions["B"].width = 80
    rows = [
        ("KOLUMNA", "ZNACZENIE / FORMAT"),
        ("ref_code", "WYMAGANE. REF/indeks produktu — taki sam jak w zamówieniu (po nim system dopasowuje artworki)."),
        ("opis_pl", "Opis po polsku."),
        ("opis_en", "Opis po angielsku."),
        ("ean", "Główny EAN-13 (13 cyfr). Opcjonalnie."),
        ("rodzina", "Rodzina/grupa produktowa. Opcjonalnie."),
        ("", ""),
        ("POZIOMY OPAKOWAŃ", "sztuka / op / opz / karton — wypełniaj tylko poziomy, które istnieją dla produktu."),
        ("{poziom}_wymiar", "Wymiary: dł x szer x wys w mm. Np. '245 x 125 x 70 mm'. Informacyjne (NIE wchodzą do porównania master↔fabryczny)."),
        ("{poziom}_ean", "EAN-13 dla danego poziomu, jeśli różni się od głównego."),
        ("{poziom}_artwork_ref", "REF artworku w systemie dla tego poziomu. Gdy nazwa pliku artworku różni się od ref_code, "
                                 "wpisz tu właściwy REF artworku (mapowanie ręczne). Puste = brak artworku na tym poziomie."),
        ("", ""),
        ("JAK UŻYĆ", "1) Jeden wiersz = jeden produkt (REF). 2) Zacznij od JEDNEJ refki testowo. 3) Zapisz jako CSV (UTF-8) lub .xlsx i zaimportuj w module Materiały."),
        ("UWAGA", "Puste komórki = pomijane. Poziomy opakowań można zmieniać w ustawieniach (sztuka/op/opz/karton to domyślne)."),
    ]
    for r, (a, b) in enumerate(rows, start=1):
        ca = ins.cell(row=r, column=1, value=a)
        cb = ins.cell(row=r, column=2, value=b)
        cb.alignment = Alignment(wrap_text=True, vertical="top")
        if r == 1 or a in ("POZIOMY OPAKOWAŃ", "JAK UŻYĆ", "UWAGA"):
            ca.font = Font(bold=True, color="1F3864"); cb.font = Font(bold=True, color="1F3864")
        else:
            ca.font = Font(bold=True)

    wb.save(out_path)
    return out_path


if __name__ == "__main__":
    out = sys.argv[1] if len(sys.argv) > 1 else "szablon_master_data.xlsx"
    print("Zapisano:", generate(out))
