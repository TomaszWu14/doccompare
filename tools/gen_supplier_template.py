#!/usr/bin/env python3
"""
tools/gen_supplier_template.py — szablon Excel do importu master daty dostawców.

Kolumny (nazwy maszynowe; komentarze w komórkach nagłówka opisują znaczenie):
  kod_dostawcy*, nazwa*, kod_producenta, kraj, waluta, incoterms,
  warunki_platnosci, lead_time_dni, slowa_kluczowe, osoba_kontaktowa,
  email, telefon, uwagi

Uruchom:  python tools/gen_supplier_template.py [out.xlsx]
Wymaga:   openpyxl
"""
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter

COLS = [
    ("kod_dostawcy", "WYMAGANE. Unikalny kod dostawcy, np. SHIELDCO.", 16),
    ("nazwa", "WYMAGANE. Pełna nazwa dostawcy.", 28),
    ("kod_producenta", "Kod producenta — TEN SAM co 'Producent' w master data materiałów (łącznik).", 18),
    ("kraj", "Kod kraju, np. CN, DE, PL.", 8),
    ("waluta", "Domyślna waluta zamówień: USD/EUR/CNY.", 10),
    ("incoterms", "Domyślne warunki dostawy, np. FOB Shanghai.", 16),
    ("warunki_platnosci", "Domyślne warunki płatności / kod SAP, np. 30 days after departure / IW04.", 22),
    ("lead_time_dni", "Czas produkcji w dniach (liczba).", 12),
    ("slowa_kluczowe", "Frazy do auto-wykrycia dostawcy z PDF, oddzielone ; lub , (np. Shieldco;FG-).", 24),
    ("osoba_kontaktowa", "Imię i nazwisko kontaktu.", 18),
    ("email", "Adres e-mail.", 20),
    ("telefon", "Numer telefonu.", 16),
    ("uwagi", "Dowolne notatki.", 26),
]

EXAMPLES = [
    {"kod_dostawcy": "SHIELDCO", "nazwa": "Shieldco Medical Co.", "kod_producenta": "OML-CHINA",
     "kraj": "CN", "waluta": "USD", "incoterms": "FOB Shanghai",
     "warunki_platnosci": "30 days after departure", "lead_time_dni": "45",
     "slowa_kluczowe": "Shieldco;FG-", "osoba_kontaktowa": "Li Wei",
     "email": "li@shieldco.cn", "telefon": "+86 21 1234567", "uwagi": "Rękawice nitrylowe"},
    {"kod_dostawcy": "EASTPORT", "nazwa": "Eastport Medical", "kod_producenta": "XN",
     "kraj": "CN", "waluta": "USD", "incoterms": "FOB Ningbo",
     "warunki_platnosci": "IZ31", "lead_time_dni": "60",
     "slowa_kluczowe": "Eastport;XN-", "osoba_kontaktowa": "Wang Fang",
     "email": "sales@eastport.cn", "telefon": "", "uwagi": ""},
]


def generate(out_path: str):
    wb = Workbook()
    ws = wb.active
    ws.title = "Dostawcy"
    hdr_font = Font(bold=True, size=10, color="1F3864")
    thin = Side(style="thin", color="B0B0B0")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for i, (name, desc, width) in enumerate(COLS, start=1):
        c = ws.cell(row=1, column=i, value=name)
        c.font = hdr_font
        c.fill = PatternFill("solid", fgColor="DDEBF7")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = border
        cm = Comment(desc, "DocCompare"); cm.width = 260; cm.height = 110
        c.comment = cm
        ws.column_dimensions[get_column_letter(i)].width = width
    name_to_idx = {n: i for i, (n, _, _) in enumerate(COLS, start=1)}
    for r, ex in enumerate(EXAMPLES, start=2):
        for name, idx in name_to_idx.items():
            cell = ws.cell(row=r, column=idx, value=ex.get(name, ""))
            cell.border = border
            cell.font = Font(size=10)
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 40

    ins = wb.create_sheet("Instrukcja")
    ins.column_dimensions["A"].width = 22
    ins.column_dimensions["B"].width = 84
    rows = [
        ("KOLUMNA", "ZNACZENIE"),
        ("kod_dostawcy", "WYMAGANE. Unikalny kod (klucz)."),
        ("nazwa", "WYMAGANE. Nazwa dostawcy."),
        ("kod_producenta", "Łącznik z master data materiałów (pole 'Producent'). Po nim wiążemy produkt↔dostawca."),
        ("slowa_kluczowe", "Auto-wykrywanie dostawcy z treści PDF. Oddzielaj ; lub ,"),
        ("warunki_platnosci/incoterms/waluta", "Domyślne wartości używane przy weryfikacji nagłówków PO/PI."),
        ("UWAGA", "Import nadpisuje dane dostawcy po 'kod_dostawcy'. Profile (mapowania kolumn) ustawiasz osobno w /suppliers."),
    ]
    for r, (a, b) in enumerate(rows, start=1):
        ca = ins.cell(row=r, column=1, value=a); cb = ins.cell(row=r, column=2, value=b)
        cb.alignment = Alignment(wrap_text=True, vertical="top")
        ca.font = Font(bold=True, color="1F3864" if r == 1 else "000000")
        if r == 1:
            cb.font = Font(bold=True, color="1F3864")
    wb.save(out_path)
    return out_path


if __name__ == "__main__":
    out = sys.argv[1] if len(sys.argv) > 1 else "szablon_master_dostawcy.xlsx"
    print("Zapisano:", generate(out))
